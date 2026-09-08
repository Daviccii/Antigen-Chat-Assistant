"""Attachment text extraction pipeline.

Phase C2: PDF / DOCX / XLSX / CSV / code text extraction. Runs as a
FastAPI BackgroundTask right after upload — attachment_routes.py calls
process_attachment(attachment_id) and returns immediately, so the
upload response doesn't block on extraction.

Image OCR/vision, audio transcription, and video are later phases —
those file_types just get parked at PROCESSED with a placeholder note
for now so they don't error out or get stuck at UPLOADED forever.
"""
import json
import logging
from pathlib import Path
from datetime import datetime
from typing import Tuple, Optional

from .attachment_models import Attachment, AttachmentType, AttachmentStatus
from .config import settings
from .db import SessionLocal

logger = logging.getLogger(__name__)

# Cap how much extracted text we ever store/inject into a prompt. This is
# a safety valve against a 300-page PDF blowing the model's context window
# — long documents get truncated with a clear marker rather than silently
# passed through in full.
MAX_EXTRACTED_CHARS = 20_000

# Spreadsheets get a structured preview, not a full dump — cap rows shown.
MAX_SPREADSHEET_PREVIEW_ROWS = 50

# Code files: read at most this many characters of source.
MAX_CODE_CHARS = 20_000


class ExtractionError(Exception):
    pass


def _truncate(text: str, limit: int = MAX_EXTRACTED_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n[... truncated, {len(text) - limit} more characters not shown ...]"


# ---------------------------------------------------------------------------
# Per-type extractors. Each returns (extracted_text, metadata_dict).
# ---------------------------------------------------------------------------

def _extract_pdf(path: Path) -> Tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages_text = []
    for page in reader.pages:
        pages_text.append(page.extract_text() or "")
    full_text = "\n\n".join(pages_text).strip()

    if not full_text:
        # No extractable text layer — likely a scanned/image-only PDF.
        # OCR fallback is a later phase; say so rather than pretending
        # we read it.
        raise ExtractionError(
            "No extractable text found — this looks like a scanned PDF. "
            "OCR support for scanned documents is coming in a later phase."
        )

    metadata = {"page_count": len(reader.pages)}
    return _truncate(full_text), metadata


def _extract_docx(path: Path) -> Tuple[str, dict]:
    import docx  # python-docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts).strip()
    if not full_text:
        raise ExtractionError("Document appears to be empty.")

    metadata = {"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)}
    return _truncate(full_text), metadata


def _extract_spreadsheet(path: Path, mime_type: str) -> Tuple[str, dict]:
    import pandas as pd

    if mime_type == "text/csv" or path.suffix.lower() == ".csv":
        df = pd.read_csv(path)
    else:
        df = pd.read_excel(path)

    total_rows, total_cols = df.shape
    preview_df = df.head(MAX_SPREADSHEET_PREVIEW_ROWS)

    summary_lines = [
        f"Spreadsheet with {total_rows} rows x {total_cols} columns.",
        f"Columns: {', '.join(str(c) for c in df.columns)}",
        "",
        f"Preview (first {min(total_rows, MAX_SPREADSHEET_PREVIEW_ROWS)} rows):",
        preview_df.to_string(index=False),
    ]

    # Basic numeric summary — genuinely useful for "analyze this data"
    # asks, and cheap to compute even on large sheets.
    numeric_cols = df.select_dtypes(include="number").columns
    if len(numeric_cols) > 0:
        summary_lines.append("")
        summary_lines.append("Numeric column summary:")
        summary_lines.append(df[numeric_cols].describe().to_string())

    full_text = "\n".join(summary_lines)
    metadata = {"rows": int(total_rows), "columns": int(total_cols), "column_names": [str(c) for c in df.columns]}
    return _truncate(full_text), metadata


def _extract_code(path: Path) -> Tuple[str, dict]:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
    except OSError as e:
        raise ExtractionError(f"Could not read file: {e}")

    if not content.strip():
        raise ExtractionError("File is empty.")

    truncated = len(content) > MAX_CODE_CHARS
    metadata = {"line_count": content.count("\n") + 1, "language_hint": path.suffix.lstrip(".")}
    return _truncate(content, MAX_CODE_CHARS), metadata


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

_EXTRACTORS = {
    AttachmentType.PDF: lambda path, mime: _extract_pdf(path),
    AttachmentType.DOCX: lambda path, mime: _extract_docx(path),
    AttachmentType.SPREADSHEET: lambda path, mime: _extract_spreadsheet(path, mime),
    AttachmentType.CODE: lambda path, mime: _extract_code(path),
}

# Types with extraction landing in later phases. Marked PROCESSED with a
# placeholder note so the UI doesn't show them stuck "processing" forever,
# but the model is told plainly that content isn't available yet.
_NOT_YET_SUPPORTED = {
    AttachmentType.IMAGE: "Image analysis (OCR/vision) isn't wired up yet — coming in a later phase.",
    AttachmentType.AUDIO: "Audio transcription isn't wired up yet — coming in a later phase.",
    AttachmentType.VIDEO: "Video analysis isn't wired up yet — coming in a later phase.",
    AttachmentType.OTHER: "This file type doesn't have an extractor yet.",
}


def process_attachment(attachment_id: int) -> None:
    """Entry point called as a BackgroundTask right after upload. Opens
    its own DB session since it runs outside the request's session scope.
    """
    db = SessionLocal()
    try:
        attachment = db.query(Attachment).filter(Attachment.id == attachment_id).first()
        if not attachment:
            logger.warning(f"process_attachment: attachment {attachment_id} not found")
            return

        attachment.status = AttachmentStatus.PROCESSING
        db.commit()

        absolute_path = Path(settings.UPLOAD_DIR) / attachment.stored_path

        if not absolute_path.exists():
            attachment.status = AttachmentStatus.FAILED
            attachment.processing_error = "File missing from storage."
            db.commit()
            return

        if attachment.file_type in _NOT_YET_SUPPORTED:
            attachment.extracted_text = _NOT_YET_SUPPORTED[attachment.file_type]
            attachment.extracted_metadata = json.dumps({"supported": False})
            attachment.status = AttachmentStatus.PROCESSED
            attachment.processed_at = datetime.utcnow()
            db.commit()
            return

        extractor = _EXTRACTORS.get(attachment.file_type)
        if not extractor:
            attachment.status = AttachmentStatus.FAILED
            attachment.processing_error = f"No extractor registered for {attachment.file_type.value}"
            db.commit()
            return

        try:
            text, metadata = extractor(absolute_path, attachment.mime_type)
        except ExtractionError as e:
            attachment.status = AttachmentStatus.FAILED
            attachment.processing_error = str(e)
            db.commit()
            logger.info(f"Attachment {attachment_id} extraction failed: {e}")
            return
        except Exception as e:
            attachment.status = AttachmentStatus.FAILED
            attachment.processing_error = "Unexpected error during extraction."
            db.commit()
            logger.exception(f"Attachment {attachment_id} extraction raised an unexpected error")
            return

        attachment.extracted_text = text
        attachment.extracted_metadata = json.dumps(metadata)
        attachment.status = AttachmentStatus.PROCESSED
        attachment.processed_at = datetime.utcnow()
        db.commit()
        logger.info(f"Attachment {attachment_id} processed OK ({attachment.file_type.value}, {len(text)} chars)")

    finally:
        db.close()


# ---------------------------------------------------------------------------
# Chat-context integration
# ---------------------------------------------------------------------------

def build_attachment_context(db, user_id: int, attachment_ids: list) -> Optional[str]:
    """Builds the block of text injected into the system prompt for a
    /chat request that references attachments. Returns None if there's
    nothing usable to add.

    Ownership is enforced here too — an attachment_id belonging to
    another user is silently skipped, not just filtered client-side.
    """
    if not attachment_ids:
        return None

    attachments = (
        db.query(Attachment)
        .filter(Attachment.id.in_(attachment_ids), Attachment.user_id == user_id)
        .all()
    )
    if not attachments:
        return None

    blocks = []
    for a in attachments:
        if a.status == AttachmentStatus.PROCESSING:
            blocks.append(f"[Attachment: {a.filename} — still being processed, not available yet]")
        elif a.status == AttachmentStatus.FAILED:
            blocks.append(f"[Attachment: {a.filename} — could not be processed: {a.processing_error}]")
        elif a.status == AttachmentStatus.PROCESSED and a.extracted_text:
            blocks.append(f"[Attachment: {a.filename} ({a.file_type.value})]\n{a.extracted_text}")

    if not blocks:
        return None

    return (
        "The user has attached the following file(s) to this message. "
        "Use their content to inform your answer:\n\n" + "\n\n---\n\n".join(blocks)
    )